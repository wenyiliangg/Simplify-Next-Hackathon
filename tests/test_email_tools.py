import io
import json
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_DIR / "backend" / "email_tools"))

import lambda_function as tools  # noqa: E402


class EmailToolsTests(unittest.TestCase):
    def event(self, **overrides):
        base = {
            "company": "Acme",
            "role": "Software Engineer Intern",
            "application_id": "A-1",
            "status": "APPLIED",
            "status_at": "2026-09-01T00:00:00Z",
            "provider": "gmail",
            "message_id": "m-1",
            "confidence": "HIGH",
            "evidence_summary": "Application receipt confirmed.",
        }
        base.update(overrides)
        return tools.normalize_event(base)

    def test_same_company_different_roles_have_different_keys(self):
        swe = self.event(role="Software Engineer Intern", application_id="")
        data = self.event(role="Data Analyst Intern", application_id="", message_id="m-2")
        self.assertNotEqual(swe["application_key"], data["application_key"])

    def test_out_of_order_events_reduce_chronologically(self):
        rejected = self.event(
            status="REJECTED", status_at="2026-09-05T00:00:00Z", message_id="m-3"
        )
        interview = self.event(
            status="INTERVIEW", status_at="2026-09-03T00:00:00Z", message_id="m-2"
        )
        applied = self.event()
        result = tools.reduce_events([rejected, applied, interview])
        self.assertEqual(result[0]["status"], "REJECTED")

    def test_as_of_reconstructs_earlier_status(self):
        applied = self.event()
        interview = self.event(
            status="INTERVIEW", status_at="2026-09-03T00:00:00Z", message_id="m-2"
        )
        result = tools.reduce_events(
            [interview, applied], as_of="2026-09-02T00:00:00Z"
        )
        self.assertEqual(result[0]["status"], "APPLIED")

    def test_formula_injection_is_escaped(self):
        self.assertEqual(tools._excel_safe("=HYPERLINK('x')"), "'=HYPERLINK('x')")
        self.assertEqual(tools._excel_safe("@SUM(A1:A2)"), "'@SUM(A1:A2)")

    def test_workbook_contains_required_sheets_and_rows(self):
        applied = self.event()
        offer = self.event(
            status="OFFER", status_at="2026-09-06T02:00:00Z", message_id="m-2"
        )
        applications = tools.reduce_events([applied, offer])
        data = tools.build_workbook(
            applications,
            [applied, offer],
            generated_at="2026-09-06T12:00:00Z",
            timezone_name="Asia/Singapore",
            provider_status={"gmail": "CONNECTED", "outlook": "NOT_CONNECTED"},
        )
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Applications", "Changes_Today", "Needs_Review"],
        )
        self.assertEqual(workbook["Applications"].max_row, 2)
        self.assertEqual(workbook["Applications"]["D2"].value, "OFFER")

    def test_synthetic_fixture_normalizes(self):
        payload = json.loads((PROJECT_DIR / "tests" / "synthetic_events.json").read_text())
        normalized = [tools.normalize_event(item) for item in payload["events"]]
        latest = tools.reduce_events(normalized)
        statuses = {item["status"] for item in latest}
        self.assertIn("REJECTED", statuses)
        self.assertIn("OFFER", statuses)
        self.assertIn("NEEDS_REVIEW", statuses)
        self.assertEqual(len(latest), 4)

    def test_production_identity_requires_trusted_lambda_context(self):
        context = type(
            "Context",
            (),
            {"client_context": type("ClientContext", (), {"custom": {"authenticatedUserId": "user-123"}})()},
        )()
        self.assertEqual(tools._resolve_user_id({}, context), "user-123")
        with self.assertRaises(PermissionError):
            tools._resolve_user_id({"demo_user_id": "attacker"}, None)


if __name__ == "__main__":
    unittest.main()
